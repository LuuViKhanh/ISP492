from docx import Document
import sys

def create_docx(path):
    doc = Document()
    
    doc.add_heading('Hướng dẫn tích hợp API Backend vào Frontend cho Operator App (Cập nhật Spec VI)', 0)
    
    doc.add_paragraph('Tài liệu này hướng dẫn chi tiết cách Frontend (FE) của Operator App tích hợp với các API Backend, được phân chia theo từng nhóm màn hình dựa trên luồng nghiệp vụ của Operator (Phiên bản Spec VI - Order-driven).')
    
    doc.add_heading('1. Màn hình Order Management (Danh sách đơn hàng)', level=1)
    
    doc.add_heading('GET /api/v1/operator/orders', level=2)
    doc.add_paragraph('Tác dụng: Lấy danh sách các Đơn hàng (Orders) từ hệ thống. Các đơn hàng này chưa được xếp vào chuyến bay nào.')
    doc.add_paragraph('Cách tích hợp: FE gọi API này để hiển thị danh sách Order. Hiển thị thông tin Origin Hub, Destination Hub, Package Info và trạng thái. Dựa vào trường planningState trả về để biết đơn hàng nào đã đến giờ cần lập lịch (READY_FOR_PLANNING).')
    
    doc.add_heading('GET /api/v1/operator/orders/{order_id}', level=2)
    doc.add_paragraph('Tác dụng: Lấy thông tin chi tiết của một Order cụ thể.')
    doc.add_paragraph('Cách tích hợp: Gọi khi Operator click vào xem chi tiết một đơn hàng trước khi bấm "Lên lịch chuyến bay".')
    
    doc.add_heading('2. Luồng Create Mission (Tạo chuyến bay)', level=1)
    
    doc.add_heading('GET /api/v1/operator/orders/{order_id}/eligible-drones', level=2)
    doc.add_paragraph('Tác dụng: Lấy danh sách các Drone khả dụng, đủ điều kiện tải trọng và hiện đang đỗ tại chính Origin Hub của đơn hàng.')
    doc.add_paragraph('Cách tích hợp: Gọi API này ở màn hình chọn thiết bị để hiển thị dropdown list các Drone cho Operator chọn. (Đã thay thế cho hàm /check-availability cũ).')
    
    doc.add_heading('POST /api/v1/operator/missions/mission-planning/analyze', level=2)
    doc.add_paragraph('Tác dụng: Phân tích lộ trình và mức độ tiêu thụ pin bằng công nghệ AI (CatBoost).')
    doc.add_paragraph('Cách tích hợp: Sau khi chọn Drone xong, FE truyền orderId và droneId vào API này. Backend sẽ trả về quãng đường (km), thời gian bay dự kiến, dự đoán năng lượng tiêu thụ (Wh) và mức độ rủi ro (Risk). Hiển thị các thông tin này lên UI.')
    doc.add_paragraph('Body Request: { "orderId": "string", "droneId": int }')
    
    doc.add_heading('POST /api/v1/operator/missions', level=2)
    doc.add_paragraph('Tác dụng: Chính thức tạo Nhiệm vụ bay (Mission).')
    doc.add_paragraph('Cách tích hợp: Khi Operator bấm "Xác nhận tạo chuyến bay", gọi API này. Backend sẽ chuyển trạng thái chuyến bay thành SCHEDULED (chờ cất cánh).')
    doc.add_paragraph('Body Request: { "orderId": "string", "droneId": int, "routeId": "string" }')
    
    doc.add_heading('3. Màn hình Hubs & Live Tracking (Không đổi)', level=1)
    doc.add_paragraph('Các API liên quan đến lấy danh sách Hubs (/operator/missions/hubs) và Live Tracking bản đồ vẫn giữ nguyên cơ chế cũ.')

    doc.save(path)
    print(f"Đã tạo file {path}")

if __name__ == "__main__":
    create_docx(r"D:\works\Đồ án kỳ 2\project\FE-Operator-APIs.docx")
